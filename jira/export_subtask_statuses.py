import os
import sys

from concurrent.futures import ThreadPoolExecutor, as_completed
from jira import JIRA
from jira.resources import Issue
from mtlogger import logger
from mtprompt import Prompt, to_list, to_bool
from openpyxl import Workbook
from openpyxl.styles import Font
from tqdm import tqdm

from _common import get_jira_client, get_jira_url

OUTPUT_DIR_PATH = os.path.join(os.path.dirname(__file__), '..', '_output', 'jira')
OUTPUT_FILE_PATH = os.path.join(OUTPUT_DIR_PATH, 'ticket_export.xlsx')
MAX_RESULTS = 255

def main():
  jira_client = get_jira_client()

  if len(sys.argv) > 1:
    ticket_ids = to_list(sys.argv[1])
    skip_completed = to_bool(sys.argv[2]) if len(sys.argv) > 2 else False
  else:
    ticket_ids = Prompt.list('Enter parent ticket IDs (comma-separated)')
    skip_completed = Prompt.bool('Skip completed tasks (Done/Closed)?')

  if not ticket_ids:
    logger.error('No ticket IDs provided.')
    return

  os.makedirs(OUTPUT_DIR_PATH, exist_ok=True)
  logger.debug(f'Processing {len(ticket_ids)} parent ticket(s)')
  export_to_excel(
    process_tickets(jira_client, ticket_ids, skip_completed)
  )
  logger.success(f'Saved "{OUTPUT_FILE_PATH}".')

def process_child(child: Issue, parent_id: str, parent_name: str, skip_completed: bool):
  child_id = child.key
  child_name = child.fields.summary
  child_status = child.fields.status.name

  if skip_completed and child_status in ('Done', 'Closed'):
    return None

  child_labels = ', '.join(getattr(child.fields, 'labels', []) or [])

  return {
    'Parent ID': parent_id,
    'Parent Name': parent_name,
    'Child ID': child_id,
    'Child Name': child_name,
    'Child Status': child_status,
    'Child Labels': child_labels,
  }

def process_tickets(jira_client: JIRA, ticket_ids: list[str], skip_completed: bool):
  data = []

  for ticket_id in tqdm(ticket_ids, desc='Processing tickets', unit='ticket'):
    ticket_id = ticket_id.strip()

    try:
      parent_issue = jira_client.issue(ticket_id)

      all_issues = []

      child_issues = parent_issue.fields.subtasks
      if child_issues:
        all_issues.extend(child_issues)

      try:
        jql_query = f'"Epic Link" = {ticket_id}'
        epic_tasks = jira_client.search_issues(jql_query, maxResults=MAX_RESULTS)
        if epic_tasks:
          existing_keys = {issue.key for issue in all_issues}
          new_epic_tasks = [issue for issue in epic_tasks if issue.key not in existing_keys]
          all_issues.extend(new_epic_tasks)
      except Exception as ex:
        logger.debug(f'No epic tasks found for {ticket_id}: {ex}')

      if not all_issues:
        continue

      tqdm.write(f'Total: {len(all_issues)} item(s) for {ticket_id}')

      parent_name = parent_issue.fields.summary

      with ThreadPoolExecutor(max_workers=5) as executor:
        futures = {executor.submit(process_child, child, ticket_id, parent_name, skip_completed): child.key for child in all_issues}

        for future in tqdm(as_completed(futures), total=len(all_issues), desc=ticket_id, leave=False, unit='child'):
          try:
            result = future.result()
            if result:
              data.append(result)
          except Exception as ex:
            logger.debug(f'Failed to process child: {ex}')

    except Exception as ex:
      logger.error(f'Failed to fetch child items for {ticket_id}: {ex}')

  return data

def export_to_excel(data: list[dict]):
  if not data:
    logger.error('No child items to export.')
    return

  try:
    jira_url = get_jira_url()
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = 'Tickets'

    fieldnames = ['Parent ID', 'Parent Name', 'Child ID', 'Child Name', 'Child Status', 'Child Labels']
    worksheet.append(fieldnames)

    for row_idx, row_data in enumerate(data, start=2):
      parent_id = row_data['Parent ID']
      parent_name = row_data['Parent Name']
      child_id = row_data['Child ID']
      child_name = row_data['Child Name']
      child_status = row_data['Child Status']
      labels = row_data['Child Labels']

      worksheet.cell(row=row_idx, column=1, value=parent_id)
      worksheet.cell(row=row_idx, column=1).hyperlink = f'{jira_url}/browse/{parent_id}'
      worksheet.cell(row=row_idx, column=1).font = Font(underline='single', color='0563C1')

      worksheet.cell(row=row_idx, column=2, value=parent_name)

      worksheet.cell(row=row_idx, column=3, value=child_id)
      worksheet.cell(row=row_idx, column=3).hyperlink = f'{jira_url}/browse/{child_id}'
      worksheet.cell(row=row_idx, column=3).font = Font(underline='single', color='0563C1')

      worksheet.cell(row=row_idx, column=4, value=child_name)
      worksheet.cell(row=row_idx, column=5, value=child_status)
      worksheet.cell(row=row_idx, column=6, value=labels)

    worksheet.column_dimensions['A'].width = 15
    worksheet.column_dimensions['B'].width = 30
    worksheet.column_dimensions['C'].width = 15
    worksheet.column_dimensions['D'].width = 30
    worksheet.column_dimensions['E'].width = 15
    worksheet.column_dimensions['F'].width = 40

    workbook.save(OUTPUT_FILE_PATH)

  except Exception as ex:
    logger.error(f'Failed to write Excel file: {ex}')

if __name__ == '__main__':
  try:
    main()
  except Exception as ex:
    logger.unhandled_error(ex)

  Prompt.enter_to_exit()
